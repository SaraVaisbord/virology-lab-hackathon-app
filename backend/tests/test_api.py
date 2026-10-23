from __future__ import annotations

import io

from openpyxl import load_workbook

from domain.vocabulary import EXPORT_COLUMN_LABELS, FIELD_LABELS, FORM_STATUS_LABELS, REASONS
from tests.conftest import make_form


def test_meta_exposes_vocabulary(client):
    meta = client.get("/meta").json()
    assert [f["key"] for f in meta["fields"]] == list(FIELD_LABELS)
    assert meta["form_statuses"] == FORM_STATUS_LABELS
    assert meta["page_size"] > 0


def test_list_is_paged_and_validated(client):
    from db import queries

    ids = [make_form(queries, file_hash=f"h{i}") for i in range(3)]
    body = client.get("/forms?limit=2").json()
    assert body["total"] == 3 and body["limit"] == 2
    assert [f["id"] for f in body["items"]] == ids[::-1][:2]

    assert client.get("/forms?limit=0").status_code == 422
    assert client.get("/forms?sort_by=nope").status_code == 400
    assert client.get("/forms?status=nope").status_code == 400


def test_confirm_field_endpoint(client):
    from db import queries

    form_id = make_form(queries)
    queries.save_extraction(form_id, [{"field_name": "id_number", "value": "1", "status": "uncertain"}])
    field_id = queries.get_form(form_id)["fields"][0]["id"]

    response = client.patch(f"/forms/{form_id}/fields/{field_id}/confirm", json={"corrected_value": " 000000018 "})
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "approved" and body["fields"][0]["value"] == "000000018"

    assert client.patch(f"/forms/{form_id}/fields/999/confirm", json={}).status_code == 404


def test_retry_and_summary(client):
    from db import queries

    form_id = make_form(queries)
    queries.set_form_status(form_id, "failed", error="x")
    assert client.post(f"/forms/{form_id}/retry").json()["status"] == "uploaded"
    assert client.get("/forms/summary").json()["pending"] == 1


def test_export_respects_filters_and_mode(client):
    from db import queries

    keep = make_form(queries, file_hash="keep")
    make_form(queries, file_hash="skip")
    queries.save_extraction(keep, [
        {"field_name": "commitment_number", "value": "1", "status": "verified"},
        {"field_name": "commitment_number", "value": "2", "status": "verified"},
    ])

    def rows(url):
        response = client.get(url)
        assert response.status_code == 200
        return list(load_workbook(io.BytesIO(response.content)).active.values)

    # One row per commitment (the row key): "keep" has two commitments → two rows
    filtered = rows("/forms/export?mode=filtered&query=keep")
    assert len(filtered) == 3 and "1" in filtered[1] and "2" in filtered[2]
    assert len(rows("/forms/export?mode=all&query=keep")) == 4  # + "skip", which has no commitment yet: one row
    assert len(rows("/forms/export?mode=approved")) == 3  # keep is approved (all fields verified)
    assert client.get("/forms/export?mode=nope").status_code == 400


def test_source_evidence_reaches_the_api(client):
    """Spec 7/8: per-source results and per-field evidence survive until the API response."""
    from db import queries
    import config
    from container import verification_rules
    from domain.verification import build_fields

    form_id = make_form(queries)
    llm = {"commitment_number": [{"value": "2646", "confidence": "high", "label": "מספר התחייבות"}]}
    barcodes = [{"data": "2646", "type": "CODE39", "x": 0.1, "y": 0.1, "width": 0.1, "height": 0.1, "page": 1}]
    queries.save_source_results(form_id, llm_result=llm, barcode_result={"status": "ok", "items": barcodes},
                                ocr_result={"status": "unavailable", "items": []})
    queries.save_extraction(form_id, build_fields(llm, barcodes, [], verification_rules(config)))

    detail = client.get(f"/forms/{form_id}").json()
    assert detail["sources"]["llm_result"] == llm
    assert detail["sources"]["barcode_result"]["items"][0]["data"] == "2646"
    assert detail["sources"]["ocr_result"]["status"] == "unavailable"
    commitment = next(f for f in detail["fields"] if f["field_name"] == "commitment_number")
    assert commitment["llm_value"] == commitment["barcode_value"] == "2646"
    assert commitment["box_source"] == "barcode"
    assert commitment["evidence"]["agreement"] == "agree"

    listed = client.get("/forms").json()["items"][0]
    assert "sources" not in listed  # heavy raw results only on the detail endpoint
    assert listed["fields"][0]["evidence"] is not None or listed["fields"][0]["status"] == "not_found"


def test_disabled_worker_is_reported_in_summary(client):
    """With WORKER_ENABLED=false the queue never starts — the UI must be told why nothing is processed."""
    from db import queries

    make_form(queries)
    summary = client.get("/forms/summary").json()
    assert summary["pending"] == 1
    assert summary["queue"] == {"running": False, "paused_until": None, "pause_reason": "disabled"}


def test_manual_add_edit_delete_without_extraction(client):
    """The lab worker can fill a form by hand, whatever its extraction state."""
    from db import queries

    form_id = make_form(queries)  # still queued — never extracted
    body = client.post(f"/forms/{form_id}/fields", json={"field_name": "commitment_number", "value": " 2646 "}).json()
    body = client.post(f"/forms/{form_id}/fields", json={"field_name": "commitment_number", "value": "2647"}).json()
    commitments = [f for f in body["fields"] if f["field_name"] == "commitment_number"]
    assert [f["value"] for f in commitments] == ["2646", "2647"]
    assert all(f["status"] == "confirmed" and f["evidence"] == {"manual": True} for f in commitments)
    assert body["status"] == "uploaded"  # a manual edit doesn't take the form out of the queue

    # single-value field: one value only — a second one must be an edit
    assert client.post(f"/forms/{form_id}/fields", json={"field_name": "id_number", "value": "000000018"}).status_code == 200
    assert client.post(f"/forms/{form_id}/fields", json={"field_name": "id_number", "value": "1"}).status_code == 400
    assert client.post(f"/forms/{form_id}/fields", json={"field_name": "nope", "value": "1"}).status_code == 400
    assert client.post(f"/forms/{form_id}/fields", json={"field_name": "sender", "value": "  "}).status_code == 400

    # edit any value (even an already confirmed one), then delete one
    field_id = commitments[1]["id"]
    edited = client.patch(f"/forms/{form_id}/fields/{field_id}/confirm", json={"corrected_value": "2648"}).json()
    assert [f["value"] for f in edited["fields"] if f["field_name"] == "commitment_number"] == ["2646", "2648"]
    after = client.delete(f"/forms/{form_id}/fields/{field_id}").json()
    assert [f["value"] for f in after["fields"] if f["field_name"] == "commitment_number"] == ["2646"]
    assert client.delete(f"/forms/{form_id}/fields/{field_id}").status_code == 404


def test_manual_value_replaces_not_found_and_completes_a_failed_form(client):
    from db import queries

    form_id = make_form(queries)
    queries.save_extraction(form_id, [{"field_name": "id_number", "status": "not_found"}])
    queries.set_form_status(form_id, "failed", error="x")
    body = client.post(f"/forms/{form_id}/fields", json={"field_name": "id_number", "value": "000000018"}).json()
    ids = [f for f in body["fields"] if f["field_name"] == "id_number"]
    assert [(f["value"], f["status"]) for f in ids] == [("000000018", "confirmed")]
    assert body["status"] == "approved"


def test_pick_a_page_number_into_a_field_keeps_its_location(client):
    from db import queries

    form_id = make_form(queries)
    queries.save_extraction(form_id, [{"field_name": "id_number", "value": "1", "llm_value": "1", "status": "uncertain"}])
    box = {"x": 0.1, "y": 0.2, "width": 0.05, "height": 0.02}

    picked = client.post(f"/forms/{form_id}/fields", json={
        "field_name": "commitment_number", "value": "947768", "box": box, "picked_from": "barcode"}).json()
    row = next(f for f in picked["fields"] if f["field_name"] == "commitment_number")
    assert (row["value"], row["status"], row["box_source"], row["x"]) == ("947768", "confirmed", "barcode", 0.1)
    assert row["evidence"] == {"manual": True, "picked_from": "barcode"}

    # single-value field: picking replaces the old value only when asked to
    assert client.post(f"/forms/{form_id}/fields", json={"field_name": "id_number", "value": "000000018"}).status_code == 400
    replaced = client.post(f"/forms/{form_id}/fields", json={
        "field_name": "id_number", "value": "000000018", "replace": True, "box": box, "picked_from": "ocr"}).json()
    ids = [f for f in replaced["fields"] if f["field_name"] == "id_number"]
    assert [(f["value"], f["box_source"]) for f in ids] == [("000000018", "ocr")]
    assert client.post(f"/forms/{form_id}/fields", json={"field_name": "sender", "value": "x", "picked_from": "nope"}).status_code == 400


def test_export_repeats_patient_details_on_every_commitment_row():
    from services.export_service import form_rows

    form = {"fields": [
        {"field_name": "commitment_number", "value": "111"},
        {"field_name": "commitment_number", "value": "222"},
        {"field_name": "central_number", "value": "9001"},
        {"field_name": "central_number", "value": "9002"},
        {"field_name": "id_number", "value": "000000018", "corrected_value": "000000026"},
    ], "duplicates": []}
    rows = form_rows(form)
    assert [row["commitment_number"] for row in rows] == ["111", "222"]
    assert all(row["central_number"] == "9001, 9002" and row["id_number"] == "000000026" for row in rows)


def test_same_commitment_on_two_forms_is_flagged_with_where_it_is(client):
    from db import queries

    first, second, other = make_form(queries, file_hash="a"), make_form(queries, file_hash="b"), make_form(queries, file_hash="c")
    box = {"x": 0.1, "y": 0.2, "width": 0.3, "height": 0.04}
    queries.save_extraction(first, [{"field_name": "commitment_number", "value": "58153", "status": "verified"}])
    queries.save_extraction(second, [{"field_name": "commitment_number", "value": "58-153", "status": "verified", "box_source": "ocr", **box}])
    queries.save_extraction(other, [{"field_name": "commitment_number", "value": "99999", "status": "verified"}])

    (dup,) = client.get(f"/forms/{first}").json()["duplicates"]  # compared as digits
    assert dup["other_form_id"] == second and dup["other_source_file"] == "b.pdf" and dup["other_page_number"] == 1
    assert (dup["x"], dup["y"], dup["box_source"]) == (0.1, 0.2, "ocr")  # where it is on the other page
    listed = {item["id"]: item["duplicates"] for item in client.get("/forms?limit=10").json()["items"]}
    assert [d["other_form_id"] for d in listed[second]] == [first] and listed[other] == []

    export = load_workbook(io.BytesIO(client.get("/forms/export?mode=all").content)).active
    header = [cell.value for cell in export[1]]
    column = header.index(EXPORT_COLUMN_LABELS["duplicate_of"])
    flagged = [row[column] for row in export.iter_rows(min_row=2, values_only=True) if row[column]]
    assert len(flagged) == 2 and any("b.pdf" in note for note in flagged)


def test_meta_gives_field_kinds_and_the_date_pattern(client):
    import re

    meta = client.get("/meta").json()
    kinds = {field["key"]: field["kind"] for field in meta["fields"]}
    assert kinds["received_date"] == "date" and kinds["commitment_number"] == "number" and kinds["id_number"] == "id"
    assert re.search(meta["date_pattern"], "20.08.26") and meta["row_key_field"] == "commitment_number"


def test_a_date_set_by_the_lab_worker_is_stored_like_extracted_dates(client):
    from db import queries

    form_id = make_form(queries)
    body = client.post(f"/forms/{form_id}/fields", json={"field_name": "received_date", "value": "20.08.26"}).json()
    (field,) = [f for f in body["fields"] if f["field_name"] == "received_date"]
    assert field["value"] == "20/08/2026"


def test_a_deleted_value_never_comes_back_on_re_extraction(client):
    """Bug seen 2026-10-21: the lab worker deleted a machine reading and re-extraction put it back."""
    from db import queries

    form_id = make_form(queries)
    machine = [
        {"field_name": "central_number", "value": "947550", "llm_value": "947550", "status": "unverified"},
        {"field_name": "central_number", "value": "123456", "llm_value": "123456", "status": "unverified"},
        {"field_name": "sender", "value": "קפלן", "llm_value": "קפלן", "status": "unverified"},
    ]
    queries.save_extraction(form_id, machine)
    fields = lambda body, name: [(f["value"], f["status"]) for f in body["fields"] if f["field_name"] == name]  # noqa: E731
    ids = {f["value"]: f["id"] for f in queries.get_form(form_id)["fields"] if f["value"]}

    client.delete(f"/forms/{form_id}/fields/{ids['947550']}")
    body = client.delete(f"/forms/{form_id}/fields/{ids['קפלן']}").json()
    assert body["rejected_values"] == {"central_number": ["947550"], "sender": ["קפלן"]}
    # the emptied field still needs a decision — the form is not silently approved without it
    assert fields(body, "sender") == [(None, "not_found")] and body["status"] == "extracted"

    queries.save_extraction(form_id, machine)  # re-extraction reads the same values again
    body = client.get(f"/forms/{form_id}").json()
    assert fields(body, "central_number") == [("123456", "unverified")]
    assert fields(body, "sender") == [(None, "not_found")]
    assert REASONS["rejected_by_worker"] in next(f["reason"] for f in body["fields"] if f["field_name"] == "sender")

    # the worker sets the value again on purpose → it is no longer rejected
    body = client.post(f"/forms/{form_id}/fields", json={"field_name": "central_number", "value": "947-550"}).json()
    assert body["rejected_values"] == {"sender": ["קפלן"]}


def test_delete_a_form_removes_its_fields_and_files(client):
    from db import queries
    from tests.conftest import png_bytes

    upload = client.post("/uploads", files={"files": ("scan.png", png_bytes(), "image/png")}).json()[0]
    (form_id,) = upload["form_ids"]
    queries.save_extraction(form_id, [{"field_name": "commitment_number", "value": "58153", "status": "verified"}])
    image_key = queries.get_form(form_id)["image_path"]
    import config

    data_dir = config.DATA_DIR
    assert (data_dir / image_key).is_file() and any((data_dir / "uploads").iterdir())

    assert client.delete(f"/forms/{form_id}").status_code == 204
    assert client.get(f"/forms/{form_id}").status_code == 404
    assert not (data_dir / image_key).exists() and not any((data_dir / "uploads").iterdir())
    assert client.get("/forms").json()["total"] == 0
    assert client.delete(f"/forms/{form_id}").status_code == 404
    # the same file can be uploaded again afterwards (it is no longer a duplicate)
    assert client.post("/uploads", files={"files": ("scan.png", png_bytes(), "image/png")}).json()[0]["status"] == "created"


def test_a_form_being_processed_cannot_be_deleted(client):
    from db import queries

    form_id = make_form(queries)
    queries.set_form_status(form_id, "processing")
    response = client.delete(f"/forms/{form_id}")
    assert response.status_code == 409 and "בעיבוד" in response.json()["detail"]
