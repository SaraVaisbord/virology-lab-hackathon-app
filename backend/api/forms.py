from __future__ import annotations

import io

from fastapi import APIRouter, HTTPException, Query, Request
from fastapi.responses import FileResponse, StreamingResponse

from db.queries import confirm_field, list_forms_with_fields
from services.forms_service import get_form, get_form_image_path, list_forms

router = APIRouter(prefix="/forms", tags=["forms"])


@router.get("/export")
@router.get("/export-file")
def export_forms_excel(
    request: Request,
    mode: str = Query(default="all"),
    status: str | None = Query(default=None),
    query: str | None = Query(default=None),
    offset: int = Query(default=0, ge=0),
    limit: int = Query(default=500, ge=1, le=5000),
):
    effective_limit = min(limit, 500)
    params = request.query_params
    raw_filters = {key: value for key, value in params.items() if key.startswith("filter_")}
    column_filters = {key.replace("filter_", ""): value for key, value in raw_filters.items()}

    rows = list_forms_with_fields(status=status, limit=effective_limit, offset=offset)

    if mode == "filtered":
        q = (query or "").strip().lower()
        rows = [
            row for row in rows
            if (not q or q in str(row.get("source_file") or "").lower() or q in str(row.get("status") or "").lower())
            and all(
                not (value and str(value).strip())
                or str(getattr(row, "get", lambda _k, default=None: default)(key, "") or "").lower().find(str(value).strip().lower()) >= 0
                for key, value in column_filters.items()
            )
        ]
    elif mode == "new":
        rows = [row for row in rows if row.get("status") == "uploaded"]

    if query:
        q = query.strip().lower()
        rows = [
            row for row in rows
            if q in str(row.get("source_file") or "").lower()
            or q in str(row.get("status") or "").lower()
            or any(
                q in str(field.get("value") or "").lower()
                or q in str(field.get("raw_value") or "").lower()
                for field in row.get("fields", [])
            )
        ]

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="openpyxl is not installed") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "Forms"
    ws.sheet_view.rightToLeft = True

    headers = [
        "מזהה",
        "שם קובץ",
        "סטטוס",
        "עמוד",
        "קוד קובץ",
        "תאריך העלאה",
        "סוג טופס",
        "שם פרטי",
        "שם משפחה",
        "מספר תעודה",
        "שולח",
        "מספר מרכזי",
        "מספר התחייבות",
    ]
    ws.append(headers)

    for row in rows:
        field_map = {field.get("field_name"): field for field in row.get("fields", [])}
        ws.append([
            row.get("id"),
            row.get("source_file"),
            row.get("status"),
            row.get("page_number"),
            row.get("file_hash"),
            row.get("uploaded_at"),
            row.get("form_type"),
            field_map.get("first_name", {}).get("value") or field_map.get("first_name", {}).get("raw_value"),
            field_map.get("last_name", {}).get("value") or field_map.get("last_name", {}).get("raw_value"),
            field_map.get("id_number", {}).get("value") or field_map.get("id_number", {}).get("raw_value"),
            field_map.get("sender", {}).get("value") or field_map.get("sender", {}).get("raw_value"),
            field_map.get("central_number", {}).get("value") or field_map.get("central_number", {}).get("raw_value"),
            field_map.get("commitment_number", {}).get("value") or field_map.get("commitment_number", {}).get("raw_value"),
        ])

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.alignment = Alignment(horizontal="right", vertical="center")

    uploaded_col_index = headers.index("תאריך העלאה") + 1
    for cell in ws.iter_cols(min_col=uploaded_col_index, max_col=uploaded_col_index, min_row=2):
        for value_cell in cell:
            if value_cell.value is None:
                continue
            if hasattr(value_cell.value, "strftime"):
                value_cell.number_format = "yyyy-mm-dd hh:mm:ss"

    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_length = 0
        for cell in column_cells:
            value = cell.value
            if value is None:
                continue
            max_length = max(max_length, len(str(value)))
        width = max(12, max_length + 2)
        if col_idx == uploaded_col_index:
            width = 22
        ws.column_dimensions[chr(64 + col_idx) if col_idx < 27 else f"{chr(64 + (col_idx // 26))}{chr(64 + (col_idx % 26))}"].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=forms_export.xlsx"},
    )


@router.get("")
def list_forms_endpoint(
    status: str | None = Query(default=None),
    limit: int = Query(default=100, ge=1, le=500),
    offset: int = Query(default=0, ge=0),
):
    rows = list_forms(status=status, limit=limit, offset=offset)
    return rows


@router.get("/{form_id}")
def get_single_form(form_id: int):
    form = get_form(form_id)
    if form is None:
        raise HTTPException(status_code=404, detail="Form not found")
    return form


@router.get("/{form_id}/image")
def get_form_image(form_id: int):
    image_path = get_form_image_path(form_id)
    if image_path is None:
        raise HTTPException(status_code=404, detail="Image not found")
    return FileResponse(path=str(image_path), media_type="image/png")


@router.patch("/{form_id}/fields/{field_name}")
def update_field_endpoint(form_id: int, field_name: str, payload: dict[str, str | None] | None = None):
    payload = payload or {}
    corrected_value = payload.get("corrected_value")
    status = payload.get("status") or "confirmed"

    try:
        result = confirm_field(form_id=form_id, field_name=field_name, corrected_value=corrected_value, status=status)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    return result
