"""Excel export — ONE ROW PER COMMITMENT (the lab's definition: the commitment number is the row key).

A form with several commitment numbers (one per item the fund approved) becomes several rows that repeat the
patient and sample details (name, ID, sender, central numbers, received date). A form with no commitment number
yet still gets one row with an empty key, so nothing uploaded silently disappears from the export.
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from db.repositories import FormFilters, FormListing
from domain.validators import digits_only
from domain.vocabulary import EXPORT_COLUMN_LABELS, EXPORT_MODE_LABELS, FIELD_LABELS, FORM_COLUMN_LABELS, FORM_STATUS_LABELS, MULTI_VALUE_FIELDS, ROW_KEY_FIELD

_MULTI_VALUE_SEPARATOR = ", "
_DATE_FORMAT = "yyyy-mm-dd hh:mm"
_MIN_COLUMN_WIDTH = 12


def _field_values(form: dict[str, Any]) -> dict[str, list[str]]:
    """Every value per field (the lab worker's correction first), each distinct value once."""
    values: dict[str, list[str]] = {}
    for field in form["fields"]:
        value = field.get("corrected_value") or field.get("value") or field.get("llm_value")
        if value and str(value) not in values.get(field["field_name"], []):
            values.setdefault(field["field_name"], []).append(str(value))
    return values


def _duplicate_note(form: dict[str, Any], commitment: str | None) -> str | None:
    """Where else this commitment number appears (it should be unique — usually the same form uploaded twice)."""
    if not commitment:
        return None
    places = [
        f'{dup["other_source_file"]} עמ\' {dup["other_page_number"]} (#{dup["other_form_id"]})'
        for dup in form.get("duplicates", []) if digits_only(dup["value"]) == digits_only(commitment)
    ]
    return _MULTI_VALUE_SEPARATOR.join(dict.fromkeys(places)) or None


def form_rows(form: dict[str, Any]) -> list[dict[str, Any]]:
    """The export rows of one form: one per commitment number, the other fields repeated on each."""
    values = _field_values(form)
    shared = {
        name: _MULTI_VALUE_SEPARATOR.join(items) if name in MULTI_VALUE_FIELDS else items[0]
        for name, items in values.items() if name != ROW_KEY_FIELD
    }
    return [
        {**shared, ROW_KEY_FIELD: commitment, "duplicate_of": _duplicate_note(form, commitment)}
        for commitment in values.get(ROW_KEY_FIELD) or [None]
    ]


def filters_for_mode(mode: str, filters: FormFilters) -> FormFilters:
    """The forms an export mode covers: "all" (every form), "approved", or the table as filtered on screen."""
    if mode not in EXPORT_MODE_LABELS:
        raise ValueError(f"Unknown export mode: {mode!r}")
    if mode == "all":
        return FormFilters(sort_by=filters.sort_by, sort_dir=filters.sort_dir)
    if mode == "approved":
        return FormFilters(status="approved", sort_by=filters.sort_by, sort_dir=filters.sort_dir)
    return filters


class ExportService:
    def __init__(self, listing: FormListing, max_rows: int) -> None:
        self._listing = listing
        self._max_rows = max_rows

    def build_forms_workbook(self, mode: str, filters: FormFilters) -> bytes:
        """The .xlsx of the forms the mode covers (filters_for_mode); ValueError for an unknown mode or filter."""
        forms, _total = self._listing.list_forms(filters_for_mode(mode, filters), limit=self._max_rows)
        return _workbook(forms)


def _workbook(forms: list[dict[str, Any]]) -> bytes:

    wb = Workbook()
    ws = wb.active
    ws.title = "Forms"
    ws.sheet_view.rightToLeft = True

    form_columns = list(FORM_COLUMN_LABELS)
    field_columns = list(FIELD_LABELS)
    extra_columns = list(EXPORT_COLUMN_LABELS)
    ws.append([FORM_COLUMN_LABELS[c] for c in form_columns] + [FIELD_LABELS[c] for c in field_columns]
              + [EXPORT_COLUMN_LABELS[c] for c in extra_columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for form in forms:
        form_cells = [FORM_STATUS_LABELS.get(form["status"], form["status"]) if c == "status" else form.get(c) for c in form_columns]
        for row in form_rows(form):
            ws.append(form_cells + [row.get(c) for c in field_columns] + [row.get(c) for c in extra_columns])

    for column_cells in ws.columns:
        width = _MIN_COLUMN_WIDTH
        for cell in column_cells:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if isinstance(cell.value, datetime):
                cell.number_format = _DATE_FORMAT
            if cell.value is not None:
                width = max(width, len(str(cell.value)) + 2)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
